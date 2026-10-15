# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FIG = Path(r"C:\Users\15736\campus-trade\homework\scope_figs")
OUT = Path(r"C:\Users\15736\Desktop\实验3-校园二手课本成色识别-前景与范围.docx")


def set_run_font(run, cn="宋体", en="Times New Roman", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def set_para_format(p, first_line=True, center=False, line=1.5, space_before=0, space_after=6, indent=None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    elif indent is not None:
        pf.first_line_indent = indent
    elif first_line:
        pf.first_line_indent = Cm(0.74)
    else:
        pf.first_line_indent = Cm(0)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def body(doc, text, **kwargs):
    p = doc.add_paragraph()
    set_para_format(p, first_line=True)
    add_text(p, text, cn="宋体", size=12, **kwargs)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=False, space_before=12, space_after=8)
    size = {1: 16, 2: 14, 3: 12}[level]
    add_text(p, text, cn="黑体", en="Times New Roman", size=size, bold=True)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=4, space_after=10, line=1.25)
    add_text(p, text, cn="宋体", size=10.5, bold=True)
    return p


def picture(doc, name, width=14.6):
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=6, space_after=2)
    run = p.add_run()
    run.add_picture(str(FIG / name), width=Cm(width))
    return p


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10.5)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "1B365D")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, bold=False, size=10.5, center=False, color=None, cn="宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, cn=cn, size=size, bold=bold, color=color)


def make_table(doc, headers, rows, col_widths, header_fill="1B365D"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate(col_widths):
        for cell in table.columns[i].cells:
            cell.width = Cm(w)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, header_fill)
        set_cell_border(cell)
        fill_cell(cell, h, bold=True, size=10.5, center=True, color=RGBColor(255, 255, 255), cn="黑体")
    for r_i, row in enumerate(rows):
        fill = "EAF2F8" if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            shade_cell(cell, fill)
            set_cell_border(cell)
            fill_cell(cell, val, size=10.5, center=(c_i == 0))
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, space_after=8)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    add_page_number(sec)

    # ===== cover =====
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=24)
    add_text(p, "北京工商大学", cn="黑体", size=22, bold=True)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=6)
    add_text(p, "计算机与人工智能学院", cn="黑体", size=16, bold=True)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=28)
    add_text(p, "实  验  报  告", cn="黑体", size=28, bold=True)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=8)
    add_text(p, "课程：软件需求工程", cn="宋体", size=14)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True)
    add_text(p, "实验名称：实验3  确定项目的前景和范围", cn="宋体", size=14)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=18)
    add_text(p, "本组题目：校园二手课本成色识别交易平台", cn="黑体", size=14, bold=True)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=6)
    add_text(p, "（图像识别判定课本新旧程度）", cn="宋体", size=12)

    info = [
        "分组：第______组",
        "组长：____________________    （角色：需求分析人员）",
        "组员：____________________    （角色：普通用户-卖家）",
        "组员：____________________    （角色：普通用户-买家）",
        "组员：____________________    （角色：高级管理者/甲方）",
        "组员：____________________    （角色：一般管理者-运营）",
        "实验时间：2026 年 10 月 ____ 日",
        "文档版本：V1.0",
    ]
    for line in info:
        p = doc.add_paragraph()
        set_para_format(p, first_line=False, center=True, space_before=4, space_after=2)
        add_text(p, line, cn="宋体", size=12)

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=28)
    add_text(p, "说明：请把姓名填在封面横线上；正文标题后已按角色标出分工，方便评分。", cn="宋体", size=10.5, color=RGBColor(91, 103, 112))

    doc.add_page_break()

    heading(doc, "一、实验内容", 1)
    body(doc, "阅读教材第5章，结合网上资料和小组讨论，参考教材第135–141页，撰写本项目的前景与范围文档。需要画图（Visio/PPT风格）和表格，用中文写。文档尽量贴近 IEEE Std 830-1998 里“先说为什么做、给谁做、做多大一圈”的写法，而不是直接跳进详细功能说明书。")
    body(doc, "本组题目选的是校园二手交易里最常见、也最适合用图像识别下手的一类东西：课本和教辅。技术重点不是再做一套闲鱼，而是：同学拍几张书的照片，系统尽量客观地判断新旧，并告诉买卖双方“为什么是这个成色”。")

    heading(doc, "二、实验要求与角色分工", 1)
    body(doc, "按组定题、定分工。报告里每个人的角色要写具体。讨论时互相当甲方和乙方。本次成绩=分组基本分+角色分。")

    make_table(
        doc,
        ["角色（评分用）", "扮演身份", "本实验具体工作（不要笼统）"],
        [
            ["需求分析人员\n（乙方，组长）", "产品/需求接口人", "主持讨论、整理业务目标与范围表、把口语改成可验收的句子、统稿与配图"],
            ["普通用户-卖家\n（甲方）", "要卖旧课本的学生", "提出拍照麻烦点、怕被砍价、怕描述不会写；确认成色等级好不好懂"],
            ["普通用户-买家\n（甲方）", "想买二手课本的学生", "提出“照片看不出内页”、怕买到涂满笔记的书；确认红框展示是否有用"],
            ["高级管理者\n（甲方）", "学院/课程项目负责人", "卡范围：本学期只做课本成色，不许把物流、支付、全品类塞进来"],
            ["一般管理者\n（甲方-运营）", "平台运营同学", "提出抽检、低置信度补拍、举报下架；确认成功指标能不能抽查"],
        ],
        [4.2, 3.6, 8.0],
    )

    heading(doc, "2.1 小组面对面讨论纪要（进入角色）", 2)
    body(doc, "讨论采用面对面方式。先由“卖家用户”和“买家用户”把真实抱怨倒出来，再由“高级管理者”一票否决做大而全，最后由需求分析人员把话翻译成目标和范围。下面是压缩后的对话要点，不是逐字稿。")
    make_table(
        doc,
        ["发言角色", "原话大意", "需求分析人员当场记下的含义"],
        [
            ["买家", "闲鱼上写九成新，拿到手全是荧光笔。", "成色必须能核对，不能只靠卖家一句话。"],
            ["卖家", "我又不是书店，让我写“轻微折痕”太专业了。", "识别结果要自动填，人只做确认和补充。"],
            ["运营", "有人故意只拍没画过的那一页怎么办？", "拍摄协议：封面+书脊+随机内页，缺一档降级。"],
            ["高管", "别跟我说要做全国物流和AI砍价。", "范围外写死：支付物流、全品类、大模型议价本期不做。"],
            ["卖方/买方", "期末大家都在卖书，系统卡了就没人用。", "成功指标里要有高峰期可用，而不是只看实验室网速。"],
        ],
        [3.0, 6.4, 6.4],
    )

    heading(doc, "三、实验结果：项目前景与范围文档", 1)
    body(doc, "以下正文按课程提供的《软件项目前景与范围文档》模板撰写，对应 IEEE 830 中与“产品前景、用户、边界、约束、验收”相关的部分。详细用例规格、接口字段放到后续《软件需求规格说明书》，本文只锁定“做不做、做到哪”。")

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=8, space_after=4)
    add_text(p, "校园二手课本成色识别交易平台", cn="黑体", size=16, bold=True)
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_after=2)
    add_text(p, "前景与范围文档（Vision and Scope）", cn="宋体", size=12)
    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_after=12)
    add_text(p, "版本 V1.0　　日期 2026-10-08　　撰写：第______组（需求分析人员统稿）", cn="宋体", size=10.5, color=RGBColor(91, 103, 112))

    heading(doc, "1. 业务需求（撰写：需求分析人员；确认：高级管理者）", 2)
    heading(doc, "1.1 背景与机遇", 3)
    body(doc, "每到期末，宿舍楼道里最常见的场景不是“有没有人要这本高数”，而是“你这本到底用成什么样了”。校园里二手课本流转本来就近、面交方便，按理说应该比闲鱼更放心。可现在同学们要么发朋友圈随手拍一张封面，要么在综合二手群里自己填成色。买家放大照片也看不清有没有折角、有没有大段笔记、纸黄不黄。卖家呢，说得好了怕被骂货不对板，说得差了又卖不出价。最后常常是面交当场翻书、当场尴尬，或者买回去才发现书被画成了“彩色练习册”。")
    body(doc, "这类矛盾，本质不是“校园里缺一个商城”，而是缺一把大家都认的尺子：这本书有多旧。尺子如果还靠人填，水分还在。本组把技术点放在图像识别上，就是想让照片自己开口：系统先找出书在画面里的位置，再标出折痕、污渍、破损、笔记、发黄这些痕迹，最后给出一个普通人能看懂的成色等级。平台仍然做发布、浏览、面约这些校园交易该有的事，但差异化就这一块——课本新旧可核对。")
    body(doc, "机遇也清楚。课本标准化程度高：几乎都是长方形、有封面有内页、磨损类型相对固定，比衣服、化妆品、自制零食更适合做视觉检测。使用场景集中在学期末和开学两周，需求密集，适合课程周期内把最小可用版本做出来做演示。学校也一直提倡闲置循环，减少资源浪费。如果我们能把“敢买、敢卖、少吵架”做成可演示的能力，题目就立得住。")
    picture(doc, "fig14_pain_chart.png", 14.4)
    caption(doc, "图1  小组访谈中最常被提到的五类抱怨（示意统计，用于支撑后面的业务目标）")
    picture(doc, "fig1_as_is_to_be.png", 14.6)
    caption(doc, "图2  现状与做成色识别之后的对比（As-Is / To-Be）")

    heading(doc, "1.2 业务目标", 3)
    body(doc, "目标按 SMART 来写：句子里要有对象、有数字、有时间，避免“提升体验”“赋能交易”这种无法验收的话。下表是给管理层看的成功定义，不是给程序员看的接口清单。")
    picture(doc, "fig15_smart.png", 14.6)
    caption(doc, "图3  业务目标写法自检：具体、可测、可做、相关、有时限")

    make_table(
        doc,
        ["编号", "业务目标描述", "成功衡量标准（本期可验收）"],
        [
            ["BO-1", "让成色从“口头承诺”变成“照片可核对”，减少货不对板引起的当面争执。", "演示集上，识别给出的等级与三人交叉人工标注相差不超过一档的比例 ≥ 80%；商品详情必须能看到瑕疵位置。"],
            ["BO-2", "让卖家定价时有参照，减少“同课本书，一个 8 元一个 28 元却看不出差别”。", "同一书名、相邻成色的在售样本，能按成色排序；发布时给出该成色的参考价区间（可先用组内样例库）。"],
            ["BO-3", "让买家更快决定买不买，缩短“反复问有没有笔记”的来回时间。", "从打开商品到发出面约，核心路径可在 2 分钟内完成；筛选条件含成色等级。"],
            ["BO-4", "降低上架成本，让“懒得写描述”的同学也愿意把书拿出来卖。", "按引导拍满规定角度后，成色和瑕疵说明自动填入，卖家平均只需确认 1 次即可提交。"],
            ["BO-5", "服务校园闲置循环，而不是做成校外电商。", "仅本校（或白名单校区）学号可交易；面交点取校园建筑，不提供快递发货入口。"],
        ],
        [2.0, 6.4, 7.4],
    )

    heading(doc, "1.3 业务风险（不做，或做砸了会怎样）", 3)
    body(doc, "风险分两头。一头是“继续现在这样”：信任成本越来越高，二手课本重新变成“宁可买新的”，闲置书最后进垃圾桶。另一头是项目本身失败：识别经常把破书判成九五新，那还不如不做，因为系统盖了章，错得更理直气壮。")
    make_table(
        doc,
        ["风险编号", "如果发生", "对业务的影响", "本期怎么挡"],
        [
            ["BR-1", "宿舍灯光差、照片模糊，识别乱给高等级", "买家更不信平台", "拍摄引导 + 清晰度检测；置信度低必须补拍，禁止直接上架高等级"],
            ["BR-2", "卖家只拍没写过的那几页", "识别被“骗”，争议回潮", "强制封面、书脊、内页；缺角度则等级封顶"],
            ["BR-3", "范围膨成全品类+支付+物流", "课程周期做不完，核心识别被挤掉", "高管角色在范围表里写死“墙外事项”"],
            ["BR-4", "同学们觉得“还不如发群里”", "没有真实试用数据", "发布流程比填表更短；识别结果要能一眼看懂"],
            ["BR-5", "期末同时上架造成排队", "关键窗口期不可用", "识别异步排队；先保证能提交，结果稍后回填"],
        ],
        [2.2, 4.6, 4.0, 5.0],
    )
    picture(doc, "fig10_risk.png", 14.4)
    caption(doc, "图4  风险矩阵：越靠右上越要先处理。光线差和只拍好面是本期最大的两颗钉子")

    heading(doc, "2. 项目前景（撰写：需求分析人员；愿景确认：全体甲方）", 2)
    heading(doc, "2.1 前景概述", 3)
    body(doc, "我们要为北京工商大学（及同类型校园试点）的同学，做一个“看得到新旧”的二手课本市场。卖家按提示拍封面、书脊和内页，图像识别给出成色等级，并在照片上用红框标出磨损和笔记；买家按成色筛选，点开就能看见系统依据，而不是只看见一句广告词。平台管身份、上架、沟通和面交预约，让书在校园里换手；识别管公平，让换手的时候少吵架。它重要，是因为课本几乎人人都有、几乎每学期都要处理，而这恰好是图像识别在校园场景里能讲清楚的一件小事。")
    picture(doc, "fig3_vision.png", 14.6)
    caption(doc, "图5  前景故事板：拍、认、判、挂、买")
    picture(doc, "fig2_context.png", 14.6)
    caption(doc, "图6  系统上下文：人和外部服务围着平台转，识别服务是独立的一方，而不是藏在按钮后面的魔术")

    heading(doc, "2.2 主要特征（特性列表）", 3)
    body(doc, "下面这些是“为实现业务目标必须具备的高层能力”。写的是用户能感知到的特性，不是类名。F-2、F-3 是技术主线，其他条目是为了让识别结果真的能进交易，而不是停留在实验报告插图里。")
    make_table(
        doc,
        ["特性编号", "系统特性（用户能看见什么）", "主要支撑的目标"],
        [
            ["F-1 拍摄引导", "发布时提示拍封面、书脊、内页；模糊、过暗、没拍到书会当场拦住。", "BO-1，BO-4"],
            ["F-2 成色自动识别", "对合格照片给出 A–E 成色、置信度；结果写回发布表单。", "BO-1，BO-4"],
            ["F-3 瑕疵红框", "在原图上标折角、污渍、破损、笔记等位置，买家能点开看。", "BO-1，BO-3"],
            ["F-4 按成色筛选", "首页和搜索可按成色、学院课程关键词、面交楼栋筛选。", "BO-3，BO-5"],
            ["F-5 参考价提示", "根据书名+成色给出组内/历史样例的价格区间，不代替卖家定价。", "BO-2"],
            ["F-6 识别留痕", "每次识别保存时间、版本、等级、框位置，供抽检和争议时回看。", "BO-1"],
            ["F-7 校园身份", "学号或校园微信登录，未通过校验不能交易。", "BO-5"],
            ["F-8 面交沟通", "站内留言和预约当面交易；不提供快递下单。", "BO-5，BO-3"],
            ["F-9 抽检与下架", "运营可抽识别结果、处理举报、下架明显造假。", "BO-1"],
            ["F-10 低置信复核", "分数没把握时要求补拍或卖家手选成色（且页面标明“卖家自报”）。", "BO-1，BO-4"],
        ],
        [3.2, 8.6, 4.0],
    )
    picture(doc, "fig6_trace.png", 14.6)
    caption(doc, "图7  业务目标到系统特性的追溯：评审时顺着箭头问“砍掉这条，目标还在不在”")
    picture(doc, "fig5_grades.png", 14.6)
    caption(doc, "图8  成色等级：人和模型用同一把尺子。A 近全新 …… E 较差")
    picture(doc, "fig4_pipeline.png", 14.6)
    caption(doc, "图9  成色识别流水线。真正的技术点在 L2 感知层：先找到书，再分割瑕疵，最后才打分")
    picture(doc, "fig9_usecase.png", 14.6)
    caption(doc, "图10  用例图（只含本学期范围内的用例）")
    picture(doc, "fig13_swimlane.png", 14.6)
    caption(doc, "图11  卖家主流程泳道：置信度不够就补拍，不把没把握的高等级写到商品上")

    heading(doc, "2.3 假设与依赖", 3)
    make_table(
        doc,
        ["类型", "条目", "如果不成立会怎样"],
        [
            ["假设", "目标用户是在校学生，会用微信拍照，能到指定楼栋面交。", "要补更重的操作说明，或根本不该做校园面交。"],
            ["假设", "课本以中文公共课/专业课纸质书为主，外形接近长方形书籍。", "活页、实验报告册、复印件识别会差，需排除或降级。"],
            ["假设", "卖家愿意按引导拍满三个角度，以换取少填表。", "F-4、F-2 的数据质量会垮，只能改回纯手填。"],
            ["依赖", "课程允许使用校园网或实验服务器存放照片（不出校演示）。", "只能本地演示，无法给宿舍同学试用。"],
            ["依赖", "微信小程序能调起摄像头并上传图片（或用H5降级方案）。", "拍摄引导整条链路要改成浏览器方案。"],
            ["依赖", "本期可用学号白名单模拟“校园身份”，不强制对接学校统一认证。", "身份真实性变弱，BO-5 只能部分满足。"],
            ["依赖", "识别模型可在组内服务器或云主机上异步运行，小程序只显示结果。", "把模型塞进手机不现实，发布会卡住。"],
        ],
        [2.2, 8.0, 5.6],
    )

    heading(doc, "3. 项目范围（撰写：一般管理者+需求分析人员；拍板：高级管理者）", 2)
    heading(doc, "3.1 范围内与范围外", 3)
    body(doc, "范围表是防膨胀的墙。墙里是本学期要交付、要验收的；墙外可以以后再议，但现在提需求一律记录为“不承诺”。")
    picture(doc, "fig7_scope.png", 14.6)
    caption(doc, "图12  范围边界图。墙外事项不是“不重要”，是“本期不做，避免拖死识别主线”")
    make_table(
        doc,
        ["范围内（In-Scope）", "范围外（Out-of-Scope）"],
        [
            ["微信小程序（或等价H5）上的课本发布、浏览、搜索、收藏", "独立 iOS/Android 原生商店应用、桌面完整电商后台（运营只需轻量抽检页）"],
            ["课本/教辅/考试资料类纸质书的成色识别与红框展示", "数码、服饰、美妆、票券、自制食品等全品类成色"],
            ["校园身份校验（可白名单）、面交点选择、站内留言预约", "真实支付、钱包、退款清分、快递物流、仓储"],
            ["识别低置信度补拍、卖家确认、运营抽检与下架", "法院级仲裁庭、保险理赔、跨校全国市场"],
            ["参考价区间（基于样例，不做智能砍价对话）", "大模型自动议价、话术生成、扫描整书OCR入库"],
            ["用户能看懂的成色说明和简单使用提示", "学校教务/财务/一卡通系统改造"],
            ["模拟或沙箱环境下的预约闭环，用于课程验收", "正式商用的等保测评与多校区运维体系"],
        ],
        [8.0, 8.0],
    )

    heading(doc, "3.2 首要用户与涉众", 3)
    picture(doc, "fig8_stakeholder.png", 13.6)
    caption(doc, "图13  涉众洋葱图：越靠里越天天用系统，越靠外越只受影响")
    make_table(
        doc,
        ["类别", "代表角色", "主要关心什么", "本期如何回应"],
        [
            ["主要用户", "卖书学生", "拍照别太烦，别把好书判太差，定价别太瞎。", "F-1、F-2、F-5、F-10"],
            ["主要用户", "买书学生", "别买到画满的书，成色筛选要快，瑕疵要能看见。", "F-3、F-4、F-8"],
            ["主要用户", "运营同学", "有人作假能抽到，下架有记录。", "F-6、F-9"],
            ["次要用户", "系统维护（本组开发）", "识别和业务分开部署，挂了能看日志。", "第4章环境约束"],
            ["影响涉众", "课程教师/学院", "题目有技术点、范围可控、能演示、能写文档。", "本文档 + V1.0 演示"],
            ["影响涉众", "学校信息化", "账号和数据别乱连正式教务，照片别随便出校。", "白名单身份 + 校内部署假设"],
            ["间接涉众", "同学圈口碑", "好用会转发，不准会群嘲。", "宁可低置信补拍，也不要装准"],
        ],
        [2.6, 3.0, 5.8, 4.4],
    )

    heading(doc, "3.3 后续版本范围及限制", 3)
    body(doc, "后续版本只作为规划登记，不纳入本期验收。高级管理者在讨论中明确：没有 V1.0 的识别准确和红框展示，后面任何“做大平台”的故事都不成立。")
    picture(doc, "fig11_roadmap.png", 14.6)
    caption(doc, "图14  版本路线。本期锁定 V1.0；V2.0 的全品类和支付明确标成不承诺")
    make_table(
        doc,
        ["版本", "计划窗口", "增加什么", "明确限制"],
        [
            ["V0.5", "课程第4–6周", "拍摄引导、手标成色、发布浏览跑通", "可以还没有自动识别，但流程要能演示"],
            ["V1.0（本期）", "课程第7–12周", "自动成色、红框、筛选、留痕、抽检", "只做课本；支付物流不做"],
            ["V1.1", "下学期（可选）", "参考价更稳、抽检工作台、复核队列", "仍不承诺跨校和真实支付"],
            ["V2.0", "未立项", "其他品类、支付、物流等", "必须另开前景文档，不得从本期范围“顺手做掉”"],
        ],
        [2.8, 3.2, 5.4, 4.4],
    )

    heading(doc, "4. 项目环境及约束（撰写：一般管理者；技术约束确认：需求分析人员）", 2)
    heading(doc, "4.1 项目属性及操作环境", 3)
    body(doc, "这是一个课程周期内的校园互联网应用：客户端以手机为主，使用场景在宿舍、教室走廊和食堂门口，网络以校园Wi-Fi和蜂窝数据为主，光线经常不好。用户一次交互很短，不耐受“上传一张图转圈一分钟”。识别因此必须做成后台任务：先收下照片，再返回成色。")
    picture(doc, "fig12_env.png", 14.6)
    caption(doc, "图15  逻辑部署：小程序负责拍和看，业务API负责交易，识别服务单独排队，数据留在约定环境")
    make_table(
        doc,
        ["属性", "本期约定"],
        [
            ["产品类型", "校园内二手课本交易工具 + 成色识别能力"],
            ["主要终端", "手机微信小程序（H5 作为拍摄受限时的退路）"],
            ["使用时段", "日常零散发布；高峰在期末停课前后两周"],
            ["数据敏感点", "学号、面交习惯、书籍照片（可能含宿舍背景）"],
            ["可用性期望", "发布和浏览优先；识别允许短暂排队，但必须有状态"],
            ["部署位置", "课程/小组服务器或校内虚拟机，演示走 HTTPS"],
        ],
        [4.0, 12.0],
    )

    heading(doc, "4.2 解决方案约束", 3)
    body(doc, "约束是必须遵守的游戏规则，不是愿望清单。下面这些如果违反，不算“实现得更酷”，算不合格。")
    make_table(
        doc,
        ["约束类型", "规则", "为什么要有这条"],
        [
            ["技术", "识别能力做成独立服务，不把重模型塞进小程序包。", "手机算力、包大小、更新节奏都扛不住。"],
            ["技术", "业务后端优先 Java Spring Boot，数据用 MySQL；缓存可用 Redis。", "与课程技术栈和校园常见部署一致，便于交接。"],
            ["技术", "前端以 uni-app/微信小程序为主。", "目标用户已经在用微信，减少再装一个App。"],
            ["数据", "演示环境中的照片默认不出校园约定网络；日志可追溯到识别版本。", "宿舍背景、姓名页可能入镜，要能删能查。"],
            ["标准", "成色只允许 A–E 五档 + 置信度，页面用中文名称，不用“0.87新”这种分数糊弄人。", "人和模型必须用同一把尺子。"],
            ["质量", "核心识别接口在单张三图任务下，排队提示可见；超时要失败得清楚，不能一直转圈。", "期末高峰宁可慢，不能死。"],
            ["交互", "识别结果必须让卖家看见再提交；系统不得在用户不知情时改价。", "避免“算法替我卖了”的纠纷。"],
        ],
        [2.4, 7.2, 6.2],
    )

    heading(doc, "4.3 项目约束", 3)
    make_table(
        doc,
        ["约束类型", "规则"],
        [
            ["时间", "与软件需求工程/后续实现课同步，V1.0 在本学期课程截止日前可演示、可答辩。"],
            ["人力", "开发与需求团队不超过本组人数（约 4–6 人），没有专职标注团队，数据集以组内采集+交叉标注为主。"],
            ["成本", "优先用已有云主机/实验机和开源模型微调，不把预算押在商业视觉API的长期调用上（课程演示可用少量调用对照）。"],
            ["技能", "图像识别是技术点，但不假设人人都会训大模型；要能讲清流水线：检测→分割/分类→规则打分。"],
            ["组织", "范围变更必须经过“高级管理者”角色点头，并改这一份前景文档的版本号，禁止口头加需求。"],
        ],
        [2.8, 13.0],
    )

    heading(doc, "4.4 运营约束", 3)
    body(doc, "系统上线（含课程演示环境）后，运行不能只靠开发同学盯着屏幕。运营约束如下：识别服务与业务服务可分开重启；每天抽检不少于一个固定数量的新发布（课程阶段可定为每日 10 条或全部，取小者）；用户举报“货不对板”时，先看识别留痕再约双方说明；明显恶意只拍好页的，降权或下架。不在校园网或授权环境之外长期存放原图。期末高峰可降级为“先上架为卖家自报成色，识别结果稍后补”，但页面必须写清楚当前是哪一种。")

    heading(doc, "5. 成功指标与验收标准（与业务目标对应）", 2)
    body(doc, "怎样才算交得出手：功能在范围内、识别不是摆设、高峰时系统还在、文档和演示能对上。下表给甲方验收用，数字宁可保守，也不写实验室里才能出现的 99%。")
    make_table(
        doc,
        ["类别", "验收标准", "对应"],
        [
            ["功能完整性", "第3.1节范围内条目均能在演示中走通；范围外功能未冒充已完成。", "范围表"],
            ["识别可用性", "在小组构建的课本样例集（建议 ≥ 80 本、覆盖五档成色）上，与三人交叉标注相差不超过一档的比例 ≥ 80%。", "BO-1"],
            ["依据可见", "抽 10 个已识别商品，至少 9 个能在详情页看到等级和至少一个瑕疵框或“未检出瑕疵”的明确说明。", "BO-1，F-3"],
            ["发布效率", "按引导拍摄的卖家，从拍完到提交，中位操作步骤显著少于纯手填成色+描述（演示对比即可）。", "BO-4"],
            ["筛选效率", "买家能用成色条件缩小列表；从打开到发起面约的演示路径 < 2 分钟。", "BO-3"],
            ["性能底线", "浏览类页面在 50 个并发量级下可响应；识别允许异步，但 30 秒内应给出完成或明确排队状态（课程环境实测）。", "BR-5"],
            ["安全与校园", "未登录或未通过身份校验不能完成交易；商品文案按平台规则拦截联系方式（与底座约定一致）。", "BO-5"],
            ["文档一致", "现场演示的功能都能在本文档特性表找到编号；现场没有的功能不得写进“已完成”。", "IEEE 830 可追溯"],
        ],
        [2.8, 10.0, 3.0],
    )

    heading(doc, "6. 附录材料", 2)
    heading(doc, "6.1 词汇表", 3)
    make_table(
        doc,
        ["术语", "人话解释"],
        [
            ["成色", "这本书看起来有多新。本期只用 A–E 五档中文名，避免“91.2%新”。"],
            ["置信度", "模型自己觉得这次判断有多稳。低了就该补拍，而不是硬给出九五新。"],
            ["瑕疵红框", "在照片上圈出折角、污渍、破损、笔记等位置，方便买家核对。"],
            ["拍摄协议", "至少拍封面、书脊、内页。少拍就降低最高可给的成色。"],
            ["范围内/外", "本期做/不做的边界。墙外不是永远不做，是现在不承诺。"],
            ["业务目标 BO", "给学校/用户看的成功定义，必须能测。"],
            ["系统特性 F", "用户能感知的能力。BO 通过 F 实现。"],
            ["面交", "校园里当面交货。本期不走快递。"],
            ["留痕", "把识别用的模型版本、时间和框存下来，以后抽检能翻出来。"],
            ["涉众", "会用到系统、或会受到影响的人，不限于点按钮的用户。"],
        ],
        [3.6, 12.2],
    )

    heading(doc, "6.2 参考资料", 3)
    refs = [
        "[1] Wiegers K, Beatty J. Software Requirements. 教材第5章及第135–141页相关内容（前景与范围）。",
        "[2] IEEE Std 830-1998, IEEE Recommended Practice for Software Requirements Specifications.",
        "[3] 本组《校园二手交易平台 Base API》设计稿（成色评估作为独立能力，不与支付物流绑死）。",
        "[4] 课程实验3模板《确定项目的前景和范围》。",
        "[5] 小组宿舍访谈纪要（2026年10月，扮演甲乙方后整理，非严格抽样调查）。",
    ]
    for r in refs:
        p = doc.add_paragraph()
        set_para_format(p, first_line=False, space_after=4)
        add_text(p, r, cn="宋体", size=12)

    heading(doc, "6.3 附图目录", 3)
    figs = [
        "图1 痛点频次条形图",
        "图2 现状/未来对比",
        "图3 SMART 目标写法",
        "图4 风险矩阵",
        "图5 前景故事板",
        "图6 系统上下文图",
        "图7 目标—特性追溯",
        "图8 成色等级标尺",
        "图9 识别流水线",
        "图10 用例图",
        "图11 泳道图",
        "图12 范围边界",
        "图13 涉众洋葱图",
        "图14 版本路线图",
        "图15 运行环境部署",
    ]
    for f in figs:
        p = doc.add_paragraph()
        set_para_format(p, first_line=False, space_after=2)
        add_text(p, f, cn="宋体", size=12)

    heading(doc, "四、实验感想", 1)
    body(doc, "以前觉得“前景与范围”就是把功能列表抄一遍，写上我们要做登录、发布、搜索。这次讨论下来才发现，最有用的其实是那张范围外的表。一说校园二手，大家很容易把闲鱼的功能全端上来：支付、物流、客服、推荐、砍价。高级管理者角色一票否决之后，题目反而清楚了——我们真正能讲的技术点，是课本新旧能不能被照片说明。")
    body(doc, "业务目标和系统特性要分开，一开始也不习惯。BO 说的是“少扯皮、好定价”，F 说的是“红框、五档成色、补拍”。没有 F，BO 就是口号；没有 BO，F 就是一堆程序员觉得酷的模块。画追溯图的时候，有两条特性对不上任何目标，后来直接删了，这比加功能更像需求工作。")
    body(doc, "图像识别很容易写成“用深度学习精准识别”。这种句子没法验收。改成“和人工标注差不超过一档、详情页必须能看到依据、没把握就补拍”，甲方才知道怎么打分。约束也不是扫兴：把模型留在独立服务里、照片不出约定环境、系统不许偷偷改价，这些是保护项目的栏杆。")
    body(doc, "不足也很明显。访谈人数少，痛点图只能当示意；成色五档是否真符合同学口头习惯，还要拿真实照片再校准；学校统一身份这学期对接不上，BO-5 只能部分兑现。这些不该藏着，应该写进假设和后续版本。文档是契约，不是作文优美程度比赛。")
    body(doc, "建议：以后实验仍然强制画“范围内/外”和“目标—特性”两张图，并且现场演示必须带着编号走。这样评分好打，组员也不容易在实现阶段偷偷把范围墙拆掉。")

    p = doc.add_paragraph()
    set_para_format(p, first_line=False, center=True, space_before=24)
    add_text(p, "—— 报告结束（请在封面填写组号与姓名）——", cn="宋体", size=10.5, color=RGBColor(91, 103, 112))

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
